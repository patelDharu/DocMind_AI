# test_system.py
"""
DocMind AI — End-to-End System Verification Test Suite
Run with:
    .\\venv\\Scripts\\python.exe test_system.py
"""

import sys
import os
import time
from pathlib import Path

# Ensure root directory is in python path
WORKSPACE_ROOT = r"D:\docmind-ai"
if WORKSPACE_ROOT not in sys.path:
    sys.path.insert(0, WORKSPACE_ROOT)

def print_header(title: str):
    print("\n" + "=" * 60)
    print(f"  {title}")
    print("=" * 60)

def print_result(name: str, success: bool, details: str = ""):
    mark = "[PASS]" if success else "[FAIL]"
    print(f"{mark} {name}")
    if details:
        print(f"       -> {details}")

total_tests = 0
passed_tests = 0

def record(name: str, success: bool, details: str = ""):
    global total_tests, passed_tests
    total_tests += 1
    if success:
        passed_tests += 1
    print_result(name, success, details)

def main():
    print_header("DocMind AI System Verification & Health Check")

    # 1. Environment & Configuration Check
    print_header("1. Environment & Configuration")
    from dotenv import load_dotenv
    load_dotenv(os.path.join(WORKSPACE_ROOT, ".env"))

    api_key = os.getenv("GEMINI_API_KEY", "")
    stripped_key = api_key.strip("'\"")
    has_key = bool(api_key and len(stripped_key) > 10)
    record("GEMINI_API_KEY present in .env", has_key, f"Key length: {len(stripped_key)}")

    model_name = os.getenv("GEMINI_MODEL", "gemini-flash-lite-latest")
    record("GEMINI_MODEL configured", bool(model_name), f"Active Model: {model_name}")

    emb_model = os.getenv("GEMINI_EMBEDDING_MODEL", "gemini-embedding-001")
    record("GEMINI_EMBEDDING_MODEL configured", bool(emb_model), f"Active Embedding: {emb_model}")

    # 2. Multi-Model Failover Cascade Test
    print_header("2. Multi-Model Failover Cascade")
    try:
        from google import genai
        from google.genai import types
        from app.core.resilience import generate_with_cascade

        client = genai.Client(api_key=api_key.strip("'\""))
        test_prompt = "Reply with 'DocMind AI Verified' and nothing else."
        ans = generate_with_cascade(
            client=client,
            contents=test_prompt,
            config=types.GenerateContentConfig(max_output_tokens=20, temperature=0.1),
        )
        record("Cascade Generation (Gemini)", bool(ans), f"Response: {ans.strip()[:60]}")
    except Exception as e:
        record("Cascade Generation (Gemini)", False, str(e))

    # 3. Native Batch Embedding Test
    print_header("3. Native Batch Embedding & Vector Engine")
    try:
        from app.core.embedder import Embedder
        embedder = Embedder()
        sample_passages = [
            "DocMind AI provides enterprise document intelligence.",
            "RRF hybrid search combines dense vector retrieval with BM25 sparse search.",
            "Trilingual support enables English, Hindi, and Gujarati processing."
        ]
        embeddings = embedder.embed_passages(sample_passages)
        valid_emb = len(embeddings) == len(sample_passages) and len(embeddings[0]) > 100
        record("Native Batch Embedding", valid_emb, f"Embedded {len(embeddings)} passages in one call (dims: {len(embeddings[0])})")
    except Exception as e:
        record("Native Batch Embedding", False, str(e))

    # 4. Universal Document Loaders Test
    print_header("4. NotebookLM-Grade Universal Ingestion")
    from app.core.loader import load_document

    # Test TXT
    txt_path = Path("app/data/test_tmp.txt")
    txt_path.parent.mkdir(parents=True, exist_ok=True)
    txt_path.write_text("Test document content for verification.", encoding="utf-8")
    try:
        txt_records = load_document(str(txt_path))
        record("Text/Markdown Loader", len(txt_records) > 0, f"Extracted {len(txt_records[0]['text'])} chars")
    except Exception as e:
        record("Text/Markdown Loader", False, str(e))
    finally:
        if txt_path.exists(): txt_path.unlink()

    # Test CSV
    csv_path = Path("app/data/test_tmp.csv")
    csv_path.write_text("Feature,Status\nOCR,Active\nRAG,Active\nStudio,Active", encoding="utf-8")
    try:
        csv_records = load_document(str(csv_path))
        record("CSV Spreadsheet Loader (Markdown Table)", len(csv_records) > 0, f"Table chars: {len(csv_records[0]['text'])}")
    except Exception as e:
        record("CSV Spreadsheet Loader", False, str(e))
    finally:
        if csv_path.exists(): csv_path.unlink()

    # Test Excel
    try:
        import pandas as pd
        excel_path = Path("app/data/test_tmp.xlsx")
        pd.DataFrame({"Plan": ["Basic", "Pro"], "Price": [10, 30]}).to_excel(excel_path, index=False)
        excel_records = load_document(str(excel_path))
        record("Excel (.xlsx) Loader", len(excel_records) > 0, f"Sheet extracted: {len(excel_records[0]['text'])} chars")
    except Exception as e:
        record("Excel (.xlsx) Loader", False, str(e))
    finally:
        if excel_path.exists():
            try:
                excel_path.unlink()
            except Exception:
                pass

    # Test PowerPoint
    try:
        from pptx import Presentation
        pptx_path = Path("app/data/test_tmp.pptx")
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = "DocMind Slide"
        prs.save(pptx_path)
        pptx_records = load_document(str(pptx_path))
        record("PowerPoint (.pptx) Loader", len(pptx_records) > 0, f"Slide extracted: {pptx_records[0]['text'][:40]}")
    except Exception as e:
        record("PowerPoint (.pptx) Loader", False, str(e))
    finally:
        if pptx_path.exists(): pptx_path.unlink()

    # 5. Document Studio (PDF & DOCX Export) Test
    print_header("5. Document Studio (Prompt-Based Augmentation & Export)")
    try:
        from app.core.doc_editor import DocumentEditor
        editor = DocumentEditor()
        sample_markdown = "# Software Agreement\n\n## 1. Scope\nDelivery of AI platform.\n\n## 2. Payment\n40% advance, 60% completion."

        # Test DOCX Export
        out_docx = Path("app/data/generated/test_verify.docx")
        out_docx.parent.mkdir(parents=True, exist_ok=True)
        editor.export_to_docx(sample_markdown, str(out_docx), title="Verification Agreement")
        record("DOCX Document Generation", out_docx.exists() and out_docx.stat().st_size > 1000, f"Size: {out_docx.stat().st_size} bytes")

        # Test PDF Export
        out_pdf = Path("app/data/generated/test_verify.pdf")
        editor.export_to_pdf(sample_markdown, str(out_pdf), title="Verification Agreement")
        record("PDF Document Generation (reportlab)", out_pdf.exists() and out_pdf.stat().st_size > 1000, f"Size: {out_pdf.stat().st_size} bytes")

    except Exception as e:
        record("Document Studio Exporter", False, str(e))
    finally:
        if out_docx.exists(): out_docx.unlink()
        if out_pdf.exists(): out_pdf.unlink()

    # 6. Backend API Server Check (if running)
    print_header("6. Backend API Server Live Status")
    try:
        import requests
        res = requests.get("http://127.0.0.1:8000/", timeout=2)
        if res.status_code == 200:
            record("FastAPI Server on port 8000", True, "Online & healthy")
        else:
            record("FastAPI Server on port 8000", False, f"Status code: {res.status_code}")
    except Exception:
        record("FastAPI Server on port 8000", False, "Offline (Start with uvicorn app.api.main:app --reload --reload-dir app --host 127.0.0.1 --port 8000)")

    # Summary
    print_header("Test Summary")
    print(f"Total Tests Executed: {total_tests}")
    print(f"Passed: {passed_tests}")
    print(f"Failed: {total_tests - passed_tests}")
    pct = (passed_tests / total_tests) * 100 if total_tests else 0
    print(f"Health Score: {pct:.1f}%")
    print("=" * 60 + "\n")

if __name__ == "__main__":
    main()
